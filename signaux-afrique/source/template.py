import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

out = sys.argv[1]
wb = Workbook()
onglets = {
    'Opportunites': ['cle', 'score', 'type_signal', 'pays', 'zone', 'produit', 'code_hs', 'mon_statut', 'annee',
                     'imports_usd', 'croissance_annuelle_pct', 'prix_usd_kg', 'donnees_anciennes', 'code_pays',
                     'date_maj', 'part_afrique_pct', 'part_cameroun_pct', 'principaux_fournisseurs',
                     'a_retenir', 'pour_ma_ferme', 'que_faire'],
    'Signaux': ['date_detection', 'priorite', 'concerne_projet', 'type', 'source', 'pays', 'produits', 'titre',
                'a_retenir', 'que_faire', 'lien', 'acheteur', 'contact', 'date_limite', 'date', 'statut', 'notes', 'id'],
    'Contacts': ['date', 'source', 'sens', 'produits', 'pays', 'quantite', 'sujet', 'extrait', 'emails',
                 'telephones', 'vigilance', 'lien_gmail', 'statut', 'prochaine_action', 'notes', 'id'],
}
statuts = '"nouveau,à creuser,contacté,en discussion,deal,sans suite"'
entete = PatternFill('solid', fgColor='0E5C55')
premier = True
for nom, colonnes in onglets.items():
    ws = wb.active if premier else wb.create_sheet()
    premier = False
    ws.title = nom
    ws.append(colonnes)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = entete
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'
    for i, col in enumerate(colonnes, start=1):
        large = col in ('titre', 'a_retenir', 'que_faire', 'extrait', 'principaux_fournisseurs', 'sujet', 'notes', 'contact')
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 48 if large else 16
    if 'statut' in colonnes:
        lettre = ws.cell(1, colonnes.index('statut') + 1).column_letter
        dv = DataValidation(type='list', formula1=statuts, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{lettre}2:{lettre}5000')
wb.save(out)
print('ok', out)
